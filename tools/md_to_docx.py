"""md_to_docx.py — convert docs/phase6_7_8_9_explained.md to a styled .docx
(tables with real column widths, code blocks, images) so LibreOffice can
export a clean PDF. Supports the subset of Markdown used in the guide:
headings, tables, fenced code, blockquotes, images, lists, bold/italic/code.
Run:  python tools/md_to_docx.py <in.md> <out.docx>
"""

import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x2E, 0x74, 0xB5)
DARK = RGBColor(0x1F, 0x38, 0x64)


def add_runs(par, text, base_bold=False):
    """Split **bold**, *italic*, `code` inline markup into runs."""
    text = text.replace("\\*", "\x00")               # protect escaped asterisks
    token = re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)", text)
    token = [t.replace("\x00", "*") for t in token]
    for part in token:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        else:
            r = par.add_run(part)
            if base_bold:
                r.bold = True


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): hex_color})
    tcPr.append(shd)


def main(md_path, docx_path):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):                       # code block
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            continue

        if line.startswith("|") and i + 1 < len(lines) and \
                set(lines[i + 1].replace("|", "").replace(" ", "")) <= {"-", ":"}:
            rows = []
            header = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            ncols = len(header)
            table = doc.add_table(rows=1 + len(rows), cols=ncols)
            table.style = "Table Grid"
            # first column modest, rest equal — real widths survive conversion
            widths = [Cm(2.6)] + [Cm(15.8 / (ncols - 1))] * (ncols - 1) \
                if ncols > 1 else [Cm(15.8)]
            table.autofit = False
            for row_cells in table.rows:
                for cell, w in zip(row_cells.cells, widths):
                    cell.width = w
            for c, txt in zip(table.rows[0].cells, header):
                c.paragraphs[0].text = ""
                add_runs(c.paragraphs[0], txt, base_bold=True)
                shade(c, "D9E2F3")
            for r_i, row in enumerate(rows, start=1):
                for c, txt in zip(table.rows[r_i].cells, row):
                    c.paragraphs[0].text = ""
                    add_runs(c.paragraphs[0], txt)
            doc.add_paragraph()
            continue

        if line.startswith("!["):                        # image
            m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            alt, src = m.group(1), m.group(2)
            path = os.path.join(os.path.dirname(md_path), src)
            if os.path.exists(path):
                doc.add_picture(path, width=Cm(15.8))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"[image missing: {src}]")
            i += 1
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            p = doc.add_heading("", level=min(level, 3))
            add_runs(p, text)
            for r in p.runs:
                r.font.color.rgb = DARK if level == 1 else ACCENT
            i += 1
            continue

        if line.startswith("> "):
            block = [line[2:]]
            while i + 1 < len(lines) and lines[i + 1].startswith("> "):
                block.append(lines[i + 1][2:])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, " ".join(b.strip() for b in block))
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            i += 1
            continue

        if re.match(r"^\d+\. ", line) or line.startswith("- "):
            style_name = "List Number" if re.match(r"^\d+\. ", line) else "List Bullet"
            p = doc.add_paragraph(style=style_name)
            add_runs(p, re.sub(r"^(\d+\.|-) ", "", line))
            i += 1
            continue

        if line.strip() == "---":
            doc.add_paragraph()
            i += 1
            continue

        if line.strip():
            # group a multi-line paragraph into one docx paragraph so inline
            # markup spanning a newline still parses and spacing stays tight
            block = [line]
            while i + 1 < len(lines):
                nxt = lines[i + 1]
                if (not nxt.strip() or nxt.startswith(("#", "```", "|", "> ",
                                                       "- ", "!["))
                        or re.match(r"^\d+\. ", nxt) or nxt.strip() == "---"):
                    break
                block.append(nxt)
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, " ".join(b.strip() for b in block))
        i += 1

    doc.save(docx_path)
    print(f"wrote {docx_path}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
